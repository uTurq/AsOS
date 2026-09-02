from __future__ import annotations

import pytest

from asos.documents.parsing import UnsupportedDocumentTypeError, parse_document


def test_parse_plain_text(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Photosynthesis converts light energy into chemical energy.")

    units = parse_document(path)
    assert len(units) == 1
    assert "Photosynthesis" in units[0].text
    assert units[0].label is None


def test_parse_markdown(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# Chapter 4\n\nBuffers resist changes in pH.")

    units = parse_document(path)
    assert len(units) == 1
    assert "Buffers" in units[0].text


def test_parse_empty_text_file_returns_no_units(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n  ")
    assert parse_document(path) == []


def test_parse_pdf(tmp_path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Exam 1 covers chapters 1 through 4.")
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="The late policy is 10 percent per day.")
    path = tmp_path / "syllabus.pdf"
    pdf.output(str(path))

    units = parse_document(path)
    assert len(units) == 2
    assert units[0].label == "page 1"
    assert "Exam 1" in units[0].text
    assert units[1].label == "page 2"
    assert "late policy" in units[1].text


def test_parse_docx(tmp_path):
    import docx

    document = docx.Document()
    for i in range(15):
        document.add_paragraph(f"Paragraph {i} about cell biology.")
    path = tmp_path / "notes.docx"
    document.save(str(path))

    units = parse_document(path)
    # 15 paragraphs grouped into batches of 10 -> 2 units
    assert len(units) == 2
    assert "Paragraph 0" in units[0].text
    assert "Paragraph 14" in units[1].text


def test_parse_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide_layout = presentation.slide_layouts[6]  # blank layout
    for i in range(2):
        slide = presentation.slides.add_slide(slide_layout)
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        textbox.text_frame.text = f"Slide {i} content about mitosis."
    path = tmp_path / "lecture.pptx"
    presentation.save(str(path))

    units = parse_document(path)
    assert len(units) == 2
    assert units[0].label == "slide 1"
    assert "mitosis" in units[0].text
    assert units[1].label == "slide 2"


def test_unsupported_file_type_raises_clear_error(tmp_path):
    path = tmp_path / "data.xyz"
    path.write_text("whatever")
    with pytest.raises(UnsupportedDocumentTypeError, match=".xyz"):
        parse_document(path)
